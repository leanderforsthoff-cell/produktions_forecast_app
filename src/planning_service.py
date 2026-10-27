"""Orchestrierung des gesamten Planungszyklus: Forecast, Produktion, MRP und Cashbedarf."""

import datetime
import io
from typing import Any, Dict, List, Optional, TypedDict
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd

from src.forecasting import generate_system_forecast
from src.inventory_math import calculate_production_needs, check_open_productions_material_coverage
from src.mrp_math import calculate_material_requirements, calculate_cash_needs
from src.utils import format_currency, format_date_deadline, format_quantity, to_date


class OpenProdShortage(TypedDict, total=False):
    production_order_id: str
    product_number: str
    product_name: str
    target_end_date: datetime.date
    material_number: str
    material_name: str
    unit: str
    required_qty: float
    available_in_time: float
    shortage_qty: float
    reason: str


class PlanningCycleResult(TypedDict):
    production_plan: pd.DataFrame
    open_prod_shortages: List[Dict[str, Any]]
    material_orders: pd.DataFrame
    material_details: pd.DataFrame
    cash_needs: pd.DataFrame


def build_initial_planning_matrix(inventory_df, history_df, target_months):
    """
    Kombiniert Bestände und historische Verkäufe mit der Systemprognose zu einer Planungsmatrix.

    Initialisiert für jeden Zielmonat die manuelle Eingabespalte 'Manuell_M{i}'
    mit dem Vorschlagswert aus 'System_M{i}'.

    Args:
        inventory_df (pd.DataFrame): Artikelbestand mit 'Artikelnummer', 'Artikelname', 'Aktueller_Bestand'.
        history_df (pd.DataFrame): Historische Verkäufe ('Artikelnummer', 'Verkaufsmonat', 'Verkaufsmenge').
        target_months (list): Geplante Zielmonate.

    Returns:
        pd.DataFrame: Planungsmatrix mit Ist-Bestand, Systemprognose und initialen manuellen Werten.
    """
    cols = ["Artikelnummer", "Artikelname", "Aktueller_Bestand"]
    for m in range(1, len(target_months) + 1):
        cols.extend([f"System_M{m}", f"Manuell_M{m}"])

    if inventory_df is None or inventory_df.empty:
        return pd.DataFrame(columns=cols)

    forecast = generate_system_forecast(inventory_df, history_df, target_months)
    df_merged = pd.merge(inventory_df, forecast, on="Artikelnummer", how="left")

    for m in range(1, len(target_months) + 1):
        sys_col = f"System_M{m}"
        man_col = f"Manuell_M{m}"
        if sys_col in df_merged.columns:
            df_merged[man_col] = df_merged[sys_col].fillna(0).astype(int)
        else:
            df_merged[man_col] = 0

    ordered_cols = [c for c in cols if c in df_merged.columns]
    extra_cols = [c for c in df_merged.columns if c not in ordered_cols]
    return df_merged[ordered_cols + extra_cols]


def load_and_prepare_data(target_months, article_list, fetch_inventory_fn, fetch_history_fn):
    """
    Lädt Bestände sowie Historie über injizierte Daten-Fetcher und baut die Planungsmatrix auf.

    Ermöglicht saubere Entkopplung von UI-Frameworks (Streamlit-Caching) und testbare Orchestrierung.
    """
    inventory = fetch_inventory_fn(article_list)
    history = fetch_history_fn(article_list)
    return build_initial_planning_matrix(inventory, history, target_months)


def run_full_planning_cycle(
    plan_df,
    target_months,
    constraints_dict,
    df_open_prod=None,
    df_cogs=None,
    df_mat_stock=None,
    df_open_po=None,
    reference_date=None,
) -> PlanningCycleResult:
    """
    Führt den gesamten Planungszyklus aus:
    1. Nettoproduktionsbedarf berechnen
    2. Engpässe offener Produktionsaufträge prüfen
    3. Materialbedarfsplanung (MRP) und Stücklistenauflösung berechnen
    4. Monatlichen Cashbedarf ermitteln
    """
    ref_date = to_date(reference_date, default=datetime.date.today())

    # 1. Produktionsbedarf
    production_plan = calculate_production_needs(
        plan_df,
        target_months,
        constraints_dict,
        df_open_prod=df_open_prod,
        df_cogs=df_cogs,
        df_mat_stock=df_mat_stock,
        df_open_po=df_open_po,
        reference_date=ref_date,
    )

    # 2. Engpässe offener Produktionen
    open_prod_shortages = check_open_productions_material_coverage(
        df_open_prod=df_open_prod,
        df_cogs=df_cogs,
        df_mat_stock=df_mat_stock,
        df_open_po=df_open_po,
        reference_date=ref_date,
    )

    # 3. MRP & Cashbedarf
    material_orders = pd.DataFrame()
    material_details = pd.DataFrame()
    if df_cogs is not None and not df_cogs.empty:
        material_orders, material_details = calculate_material_requirements(
            production_plan,
            df_cogs,
            df_mat_stock if df_mat_stock is not None else pd.DataFrame(),
            df_open_po=df_open_po,
            reference_date=ref_date,
            df_open_prod=df_open_prod,
        )

    cash_needs = calculate_cash_needs(material_orders, reference_date=ref_date)

    return {
        "production_plan": production_plan,
        "open_prod_shortages": open_prod_shortages,
        "material_orders": material_orders,
        "material_details": material_details,
        "cash_needs": cash_needs,
    }


def extract_out_of_stock_warnings(production_plan, month_labels):
    """Extrahiert lesbare Warnmeldungen für künftige Fehlbestände aus dem Produktionsplan."""
    warnings = []
    for m_idx, m_label in enumerate(month_labels, start=1):
        s_col, r_col = f"Fehlbestand_M{m_idx}", f"Fehlbestand_Grund_M{m_idx}"
        if s_col in production_plan.columns:
            for aff in production_plan[production_plan[s_col] > 0].itertuples(index=False):
                formatted_shortage = format_quantity(getattr(aff, s_col))
                raw_r = getattr(aff, r_col, None)
                reason = (
                    str(raw_r)
                    if pd.notna(raw_r) and str(raw_r).strip()
                    else "Bestelldeadline ist abgelaufen – Produktion nicht mehr rechtzeitig möglich!"
                )
                art_name = getattr(aff, "Artikelname", getattr(aff, "Artikelnummer", "Artikel"))
                warnings.append(
                    f"**{art_name}**: Drohende Fehlmenge von **{formatted_shortage} Stück** im **{m_label}** ({reason})"
                )
    return warnings


def format_display_plan(production_plan, target_month_dates, month_labels):
    """Bereitet den berechneten Produktionsplan für die tabellarische UI-Anzeige auf."""
    display_plan = production_plan.copy()
    for m in range(1, len(target_month_dates) + 1):
        order_col = f"Bestelldatum_M{m}"
        shortage_col = f"Fehlbestand_M{m}"
        prod_col = f"Produktion_M{m}"
        sched_col = f"Angesetzte_Produktion_M{m}"

        if order_col in display_plan.columns:
            display_plan[order_col] = display_plan[order_col].apply(format_date_deadline)
        if sched_col in display_plan.columns:
            display_plan[sched_col] = display_plan[sched_col].apply(format_quantity)
        if shortage_col in display_plan.columns and prod_col in display_plan.columns:
            display_plan[prod_col] = np.where(
                display_plan[shortage_col] > 0,
                display_plan[prod_col].astype(str) + " ⚠️",
                display_plan[prod_col].astype(str),
            )

    ui_columns = ["Artikelname", "MOQ", "Lead_Time_Weeks"]
    rename_map = {"Lead_Time_Weeks": "Vorlauf (Wochen)"}

    for m_idx, m_label in enumerate(month_labels, start=1):
        ui_columns.extend([f"Angesetzte_Produktion_M{m_idx}", f"Produktion_M{m_idx}", f"Bestelldatum_M{m_idx}"])
        rename_map[f"Angesetzte_Produktion_M{m_idx}"] = f"Angesetzt {m_label}"
        rename_map[f"Produktion_M{m_idx}"] = f"Produktion {m_label}"
        rename_map[f"Bestelldatum_M{m_idx}"] = f"Order für {m_label}"

    valid_ui_cols = [c for c in ui_columns if c in display_plan.columns]
    return display_plan[valid_ui_cols].rename(columns=rename_map)


def prepare_supplier_order_summary(sup_df: pd.DataFrame) -> pd.DataFrame:
    """
    Erstellt eine chronologisch nach Produktionsdatum aufsteigend sortierte
    Zusammenfassung des Bestellvolumens pro Produktion für einen Lieferanten.

    Args:
        sup_df (pd.DataFrame): Bestellungen eines Lieferanten mit 'Für_Produktion_Am'
                               und 'Gesamtpreis (€)' (bzw. 'Gesamtpreis').

    Returns:
        pd.DataFrame: Aggregiertes Bestellvolumen pro Produktion, aufsteigend sortiert nach
                      Produktionsdatum. Enthält 'Für_Produktion_Am', 'sum', 'count',
                      'formatted_cost', 'delta_text'.
    """
    cols = ["Für_Produktion_Am", "sum", "count", "formatted_cost", "delta_text"]
    if sup_df is None or sup_df.empty:
        return pd.DataFrame(columns=cols)

    price_col = "Gesamtpreis (€)" if "Gesamtpreis (€)" in sup_df.columns else "Gesamtpreis"
    if price_col not in sup_df.columns or "Für_Produktion_Am" not in sup_df.columns:
        return pd.DataFrame(columns=cols)

    df = sup_df.copy()
    if "Produktion_Datum" in df.columns:
        df["_sort_date"] = df["Produktion_Datum"].apply(to_date)
    else:
        df["_sort_date"] = df["Für_Produktion_Am"].apply(to_date)

    min_date = datetime.date.min
    df["_sort_date"] = df["_sort_date"].apply(lambda d: d if d is not None else min_date)

    summary = (
        df.groupby(["_sort_date", "Für_Produktion_Am"], dropna=False)[price_col]
        .agg(["sum", "count"])
        .reset_index()
    )
    summary = summary.sort_values(by="_sort_date", ascending=True).reset_index(drop=True)
    summary["formatted_cost"] = summary["sum"].apply(format_currency)
    summary["delta_text"] = summary["count"].astype(int).apply(
        lambda cnt: f"{cnt} Position{'en' if cnt != 1 else ''}"
    )
    return summary


def prepare_material_orders_export(material_orders: Optional[pd.DataFrame]) -> pd.DataFrame:
    """
    Bereitet die Materialbestellungen für den Excel-Export vor.

    Sortiert zwingend nach Lieferant, gefolgt vom spätesten Bestelldatum,
    und stellt sicher, dass Einzelpreis (€) und Gesamtpreis (€) pro Position ausgewiesen werden.
    """
    export_cols = [
        "Lieferant",
        "Material_Nr",
        "Material_Name",
        "Bestellmenge",
        "Einheit",
        "Einzelpreis (€)",
        "Gesamtpreis (€)",
        "Vorlaufzeit_Tage",
        "Spätestes_Bestelldatum",
        "Für_Produktion_Am",
        "Benötigt_Für_Produkt",
    ]
    if material_orders is None or material_orders.empty:
        return pd.DataFrame(columns=export_cols)

    df = material_orders.copy()

    # Preisfelder vereinheitlichen
    if "Einzelpreis (€)" not in df.columns and "Einzelpreis" in df.columns:
        df["Einzelpreis (€)"] = df["Einzelpreis"]
    if "Gesamtpreis (€)" not in df.columns and "Gesamtpreis" in df.columns:
        df["Gesamtpreis (€)"] = df["Gesamtpreis"]

    # Datumsfelder säubern
    if "Spätestes_Bestelldatum" in df.columns:
        df["Spätestes_Bestelldatum"] = df["Spätestes_Bestelldatum"].apply(to_date)
    if "Für_Produktion_Am" in df.columns:
        df["Für_Produktion_Am"] = df["Für_Produktion_Am"].apply(to_date)

    # Numerische Typen sicherstellen
    if "Bestellmenge" in df.columns:
        df["Bestellmenge"] = pd.to_numeric(df["Bestellmenge"], errors="coerce").fillna(0)
    if "Einzelpreis (€)" in df.columns:
        df["Einzelpreis (€)"] = pd.to_numeric(df["Einzelpreis (€)"], errors="coerce").fillna(0.0).round(2)
    if "Gesamtpreis (€)" in df.columns:
        df["Gesamtpreis (€)"] = pd.to_numeric(df["Gesamtpreis (€)"], errors="coerce").fillna(0.0).round(2)
    if "Vorlaufzeit_Tage" in df.columns:
        df["Vorlaufzeit_Tage"] = pd.to_numeric(df["Vorlaufzeit_Tage"], errors="coerce").fillna(0).astype(int)

    # Nach Lieferant (aufsteigend) und spätestem Bestelldatum sortieren
    sort_cols = [c for c in ["Lieferant", "Spätestes_Bestelldatum", "Material_Nr"] if c in df.columns]
    if sort_cols:
        df = df.sort_values(by=sort_cols, ascending=True, na_position="last").reset_index(drop=True)

    # Fehlende Spalten auffüllen falls nötig
    for c in export_cols:
        if c not in df.columns:
            df[c] = None

    return df[export_cols]


def prepare_production_orders_export(
    production_plan: Optional[pd.DataFrame],
    target_month_dates: List[Any],
    month_labels: List[str],
) -> pd.DataFrame:
    """
    Bereitet die anstehenden Produktionsaufträge (Menge > 0) für den Excel-Export vor.

    Erstellt eine chronologische Auftragsliste, aus der hervorgeht, wann welche Menge
    eines Fertigprodukts beauftragt werden muss und zu wann die Fertigstellung geplant ist.
    """
    export_cols = [
        "Spätestes_Startdatum",
        "Ziel_Fertigstellung",
        "Zielmonat",
        "Artikelnummer",
        "Artikelname",
        "Produktionsmenge",
        "Vorlaufzeit_Wochen",
        "Mindestbestellmenge (MOQ)",
    ]
    if production_plan is None or production_plan.empty or not target_month_dates:
        return pd.DataFrame(columns=export_cols)

    rows = []
    num_months = min(len(target_month_dates), len(month_labels))

    for _, row in production_plan.iterrows():
        art_nr = str(row.get("Artikelnummer", "")).strip()
        art_name = str(row.get("Artikelname", "")).strip()
        lt_w = int(row.get("Lead_Time_Weeks", 0)) if pd.notna(row.get("Lead_Time_Weeks")) else 0
        moq_val = int(row.get("MOQ", 0)) if pd.notna(row.get("MOQ")) else 0

        for m_idx in range(1, num_months + 1):
            prod_col = f"Produktion_M{m_idx}"
            order_col = f"Bestelldatum_M{m_idx}"
            qty_raw = row.get(prod_col, 0)
            qty = float(qty_raw) if pd.notna(qty_raw) else 0.0

            if qty > 0:
                rows.append({
                    "Spätestes_Startdatum": to_date(row.get(order_col)),
                    "Ziel_Fertigstellung": to_date(target_month_dates[m_idx - 1]),
                    "Zielmonat": month_labels[m_idx - 1],
                    "Artikelnummer": art_nr,
                    "Artikelname": art_name,
                    "Produktionsmenge": int(qty),
                    "Vorlaufzeit_Wochen": lt_w,
                    "Mindestbestellmenge (MOQ)": moq_val,
                })

    if not rows:
        return pd.DataFrame(columns=export_cols)

    df_res = pd.DataFrame(rows)
    df_res = df_res.sort_values(
        by=["Spätestes_Startdatum", "Artikelname"],
        ascending=True,
        na_position="last",
    ).reset_index(drop=True)
    return df_res[export_cols]


def export_planning_to_excel(
    material_orders: Optional[pd.DataFrame],
    production_plan: Optional[pd.DataFrame],
    target_month_dates: List[Any],
    month_labels: List[str],
) -> bytes:
    """
    Erstellt eine Excel-Arbeitsmappe (.xlsx) mit zwei formatierten Sheets:
    1. 'Materialbestellungen': Sortiert nach Lieferant mit Einzel- und Gesamtpreisen.
    2. 'Produktionsaufträge': Anstehende Produktionen mit Beauftragungs- und Fertigstellungsterminen.

    Gibt den Dateiinhalt als Bytes für den Streamlit-Download zurück.
    """
    df_mat = prepare_material_orders_export(material_orders)
    df_prod = prepare_production_orders_export(production_plan, target_month_dates, month_labels)

    wb = openpyxl.Workbook()

    # Sheet 1: Materialbestellungen
    ws_mat = wb.active
    ws_mat.title = "Materialbestellungen"

    # Sheet 2: Produktionsaufträge
    ws_prod = wb.create_sheet(title="Produktionsaufträge")

    header_fill = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
    header_font = Font(bold=True, size=11)

    sheets_and_data = [
        (
            ws_mat,
            df_mat,
            {
                "Einzelpreis (€)": "#,##0.00 €",
                "Gesamtpreis (€)": "#,##0.00 €",
                "Bestellmenge": "#,##0",
                "Vorlaufzeit_Tage": "#,##0",
                "Spätestes_Bestelldatum": "DD.MM.YYYY",
                "Für_Produktion_Am": "DD.MM.YYYY",
            },
        ),
        (
            ws_prod,
            df_prod,
            {
                "Produktionsmenge": "#,##0",
                "Vorlaufzeit_Wochen": "#,##0",
                "Mindestbestellmenge (MOQ)": "#,##0",
                "Spätestes_Startdatum": "DD.MM.YYYY",
                "Ziel_Fertigstellung": "DD.MM.YYYY",
            },
        ),
    ]

    for ws, df, num_formats in sheets_and_data:
        # Kopfzeile schreiben
        ws.append(list(df.columns))
        ws.freeze_panes = "A2"

        # Kopfzeilen-Styling
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center" if "datum" in str(cell.value).lower() else "left")

        # Datenzeilen schreiben
        for row_vals in df.itertuples(index=False):
            ws.append(list(row_vals))

        # Zell-Formatierungen anwenden
        col_names = list(df.columns)
        for row_idx in range(2, ws.max_row + 1):
            for col_idx, col_name in enumerate(col_names, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if col_name in num_formats and cell.value is not None:
                    cell.number_format = num_formats[col_name]
                    if "€" in num_formats[col_name] or col_name in {"Bestellmenge", "Produktionsmenge"}:
                        cell.alignment = Alignment(horizontal="right")
                    elif "DD.MM.YYYY" in num_formats[col_name]:
                        cell.alignment = Alignment(horizontal="center")

        # Spaltenbreiten automatisch anpassen
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val_str = ""
                if cell.value is not None:
                    if isinstance(cell.value, (datetime.date, datetime.datetime)):
                        val_str = cell.value.strftime("%d.%m.%Y")
                    elif isinstance(cell.value, float):
                        val_str = f"{cell.value:,.2f}"
                    else:
                        val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


